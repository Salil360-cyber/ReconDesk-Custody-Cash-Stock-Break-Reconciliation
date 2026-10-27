"""
Excel daily control report (the file an ops analyst signs off).

Sheets
------
Sign-off     KPIs, break counts by type, and a sign-off checklist. All figures
             are live Excel formulas pointing at the other sheets, so if the
             analyst updates an owner or status, the checklist updates too.
Breaks       Every open break (cash + stock) with aging, priority, cause and
             next action. Yellow columns are for the analyst to fill in.
Aging        Break count and value by age bucket and priority.
Cash Recon   Every cash row, matched or not (full audit trail).
Stock Recon  Every position row, matched or not.
Assumptions  Every tolerance, threshold and rate the engine used.
"""

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3A5F")
BOLD = Font(name=FONT, bold=True)
NORMAL = Font(name=FONT)
WHITE_BOLD = Font(name=FONT, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = '#,##0.00;(#,##0.00);"-"'
INT = '#,##0;(#,##0);"-"'

CASH_STATUSES = ["AMOUNT_DIFF", "TIMING_DIFF", "MISSING_AT_CUSTODIAN", "MISSING_IN_BOOKS",
                 "DUPLICATE_IN_BOOKS", "DUPLICATE_AT_CUSTODIAN"]
STOCK_STATUSES = ["QTY_BREAK", "MISSING_AT_CUSTODIAN", "MISSING_IN_BOOKS"]

BREAK_COLS = [
    ("Break ID", 11, "break_id"), ("Recon", 8, "recon_type"), ("Account", 11, "account_id"),
    ("Reference / Position", 24, "key"), ("Txn Type / Security", 30, "item"),
    ("Ccy", 6, "currency"), ("Break Type", 24, "status"),
    ("Books", 15, "books_value"), ("Custodian", 15, "custodian_value"),
    ("Difference", 15, "difference"), ("Difference (INR eq.)", 18, "difference_inr"),
    ("Books Date", 12, "books_value_date"), ("Custodian Date", 13, "custodian_value_date"),
    ("Age (days)", 10, "age_days"), ("Age Bucket", 11, "age_bucket"), ("Priority", 9, "priority"),
    ("Probable Cause", 28, "probable_cause"), ("Explanation", 55, "explanation"),
    ("Next Action", 60, "next_action"),
    ("Owner", 14, None), ("Resolution Status", 17, None), ("Comments", 30, None),
]

RECON_COLS = [
    ("Recon", 8, "recon_type"), ("Account", 11, "account_id"), ("Reference / Position", 24, "key"),
    ("Txn Type / Security", 30, "item"), ("Ccy", 6, "currency"), ("Books", 15, "books_value"),
    ("Custodian", 15, "custodian_value"), ("Status", 26, "status"),
    ("Difference", 15, "difference"), ("Books Date", 12, "books_value_date"),
    ("Custodian Date", 13, "custodian_value_date"),
]


def _header(ws, row, cols):
    for i, (name, width, _) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.border = WHITE_BOLD, HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30


def _write_table(ws, cols, rows, number_cols=()):
    _header(ws, 1, cols)
    for r_i, r in enumerate(rows, start=2):
        for c_i, (_, _, field) in enumerate(cols, start=1):
            val = r.get(field, "") if field else ""
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font, cell.border = NORMAL, BORDER
            if c_i in number_cols and isinstance(val, (int, float)):
                cell.number_format = NUM if isinstance(val, float) else INT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}"


def build_report(path, cash_results, stock_results, breaks, as_of, ai_mode):
    wb = Workbook()
    signoff = wb.active
    signoff.title = "Sign-off"
    ws_b = wb.create_sheet("Breaks")
    ws_a = wb.create_sheet("Aging")
    ws_c = wb.create_sheet("Cash Recon")
    ws_s = wb.create_sheet("Stock Recon")
    ws_x = wb.create_sheet("Assumptions")

    # ---------------- Breaks ----------------
    for i, b in enumerate(breaks, start=1):
        b["break_id"] = f"BRK-{i:04d}"
    _write_table(ws_b, BREAK_COLS, breaks, number_cols={8, 9, 10, 11})
    n = len(breaks)
    last = n + 1
    for r in range(2, last + 1):
        for col in (20, 21, 22):
            ws_b.cell(row=r, column=col).fill = INPUT_FILL
        ws_b.cell(row=r, column=21, value="OPEN")
        for col in (18, 19):
            ws_b.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"OPEN,IN PROGRESS,ESCALATED,RESOLVED"', allow_blank=False)
    ws_b.add_data_validation(dv)
    dv.add(f"U2:U{last}")
    red = PatternFill("solid", fgColor="F8CBAD")
    amber = PatternFill("solid", fgColor="FFE699")
    ws_b.conditional_formatting.add(f"P2:P{last}", CellIsRule(operator="equal", formula=['"HIGH"'], fill=red))
    ws_b.conditional_formatting.add(f"P2:P{last}", CellIsRule(operator="equal", formula=['"MEDIUM"'], fill=amber))

    # ---------------- Cash / Stock recon ----------------
    _write_table(ws_c, RECON_COLS, cash_results, number_cols={6, 7, 9})
    _write_table(ws_s, RECON_COLS, stock_results, number_cols={6, 7, 9})
    nc, ns = len(cash_results) + 1, len(stock_results) + 1

    # Ranges on the Breaks sheet
    R = lambda col: f"Breaks!${col}$2:${col}${last}"  # noqa: E731
    REC, TYPE, DIFF_INR, BUCKET, PRIO, OWNER, RES = (R("B"), R("G"), R("K"), R("O"), R("P"), R("T"), R("U"))

    # ---------------- Sign-off ----------------
    ws = signoff
    ws.column_dimensions["A"].width = 44
    for col in "BCD":
        ws.column_dimensions[col].width = 18
    ws["A1"] = "ReconDesk - Daily Custody Reconciliation Control"
    ws["A1"].font = TITLE
    ws["A2"] = f"Business date (as of): {as_of.isoformat()}"
    ws["A3"] = f"Explanations generated by: {ai_mode}"
    ws["A4"] = "Synthetic data for portfolio demonstration. Figures below are live formulas."
    for c in ("A2", "A3", "A4"):
        ws[c].font = Font(name=FONT, italic=True, color="595959")

    def label_row(row, label, *cells, fmt=None, bold=False):
        ws.cell(row=row, column=1, value=label).font = BOLD if bold else NORMAL
        for i, v in enumerate(cells, start=2):
            c = ws.cell(row=row, column=i, value=v)
            c.font = BOLD if bold else NORMAL
            c.border = BORDER
            if fmt:
                c.number_format = fmt

    def section(row, title, *heads):
        ws.cell(row=row, column=1, value=title).font = WHITE_BOLD
        ws.cell(row=row, column=1).fill = HEADER_FILL
        for i, h in enumerate(heads, start=2):
            c = ws.cell(row=row, column=i, value=h)
            c.font, c.fill = WHITE_BOLD, HEADER_FILL
            c.alignment = Alignment(horizontal="center")

    section(6, "Headline", "Cash", "Stock")
    CS = f"'Cash Recon'!$H$2:$H${nc}"
    SS = f"'Stock Recon'!$H$2:$H${ns}"
    label_row(7, "Items reconciled", f"=COUNTA({CS})", f"=COUNTA({SS})", fmt=INT)
    label_row(8, "Matched (incl. within tolerance)",
              f'=COUNTIF({CS},"MATCHED")+COUNTIF({CS},"MATCHED_WITHIN_TOLERANCE")',
              f'=COUNTIF({SS},"MATCHED")', fmt=INT)
    label_row(9, "Breaks", "=B7-B8", "=C7-C8", fmt=INT)
    label_row(10, "Match rate", "=IF(B7=0,0,B8/B7)", "=IF(C7=0,0,C8/C7)", fmt="0.0%", bold=True)
    label_row(11, "Break value (INR eq., absolute)",
              f'=SUMPRODUCT(({REC}="CASH")*ABS({DIFF_INR}))',
              f'=SUMPRODUCT(({REC}="STOCK")*ABS({DIFF_INR}))', fmt=NUM)
    label_row(12, "HIGH priority breaks",
              f'=COUNTIFS({REC},"CASH",{PRIO},"HIGH")',
              f'=COUNTIFS({REC},"STOCK",{PRIO},"HIGH")', fmt=INT)

    section(14, "Cash breaks by type", "Count", "Value (INR eq.)")
    row = 15
    for s in CASH_STATUSES:
        label_row(row, s, f'=COUNTIFS({REC},"CASH",{TYPE},"{s}")',
                  f'=SUMPRODUCT(({REC}="CASH")*({TYPE}="{s}")*ABS({DIFF_INR}))')
        ws.cell(row=row, column=2).number_format = INT
        ws.cell(row=row, column=3).number_format = NUM
        row += 1
    section(row + 1, "Stock breaks by type", "Count", "Value (INR eq.)")
    row += 2
    for s in STOCK_STATUSES:
        label_row(row, s, f'=COUNTIFS({REC},"STOCK",{TYPE},"{s}")',
                  f'=SUMPRODUCT(({REC}="STOCK")*({TYPE}="{s}")*ABS({DIFF_INR}))')
        ws.cell(row=row, column=2).number_format = INT
        ws.cell(row=row, column=3).number_format = NUM
        row += 1

    row += 1
    section(row, "Sign-off checklist", "Result")
    checks = [
        ("Cash reconciliation completed", '=IF(B7>0,"YES","NO")'),
        ("Stock reconciliation completed", '=IF(C7>0,"YES","NO")'),
        ("Every HIGH break has an owner",
         f'=IF(COUNTIFS({PRIO},"HIGH",{OWNER},"")=0,"YES","NO")'),
        ("No break aged 8+ days", f'=IF(COUNTIF({BUCKET},"8+ days")=0,"YES","NO")'),
        ("Breaks still OPEN (not yet worked)", f'=COUNTIF({RES},"OPEN")'),
        ("Breaks ESCALATED", f'=COUNTIF({RES},"ESCALATED")'),
        ("Breaks RESOLVED today", f'=COUNTIF({RES},"RESOLVED")'),
    ]
    first_check = row + 1
    for label, f in checks:
        row += 1
        label_row(row, label, f)
    green = PatternFill("solid", fgColor="C6EFCE")
    ws.conditional_formatting.add(f"B{first_check}:B{row}", CellIsRule(operator="equal", formula=['"YES"'], fill=green))
    ws.conditional_formatting.add(f"B{first_check}:B{row}", CellIsRule(operator="equal", formula=['"NO"'], fill=red))

    row += 2
    for label in ("Prepared by", "Reviewed by", "Sign-off time"):
        ws.cell(row=row, column=1, value=label).font = BOLD
        c = ws.cell(row=row, column=2)
        c.fill, c.border = INPUT_FILL, BORDER
        row += 1
    ws.cell(row=row + 1, column=1,
            value="Yellow cells (here and on the Breaks sheet) are for the analyst to fill in.").font = \
        Font(name=FONT, italic=True, color="595959")

    # ---------------- Aging ----------------
    ws_a.column_dimensions["A"].width = 14
    for col in "BCDEF":
        ws_a.column_dimensions[col].width = 16
    ws_a["A1"] = "Break aging (cash + stock)"
    ws_a["A1"].font = TITLE
    heads = ["Age bucket", "HIGH", "MEDIUM", "LOW", "Total", "Value (INR eq.)"]
    for i, h in enumerate(heads, start=1):
        c = ws_a.cell(row=3, column=i, value=h)
        c.font, c.fill, c.border = WHITE_BOLD, HEADER_FILL, BORDER
    for r_i, (_, bucket) in enumerate(config.AGE_BUCKETS, start=4):
        ws_a.cell(row=r_i, column=1, value=bucket).font = NORMAL
        for c_i, p in enumerate(["HIGH", "MEDIUM", "LOW"], start=2):
            ws_a.cell(row=r_i, column=c_i, value=f'=COUNTIFS({BUCKET},$A{r_i},{PRIO},"{p}")')
        ws_a.cell(row=r_i, column=5, value=f"=SUM(B{r_i}:D{r_i})")
        ws_a.cell(row=r_i, column=6, value=f"=SUMPRODUCT(({BUCKET}=$A{r_i})*ABS({DIFF_INR}))")
    tot = 4 + len(config.AGE_BUCKETS)
    ws_a.cell(row=tot, column=1, value="Total").font = BOLD
    for c_i in range(2, 7):
        L = get_column_letter(c_i)
        ws_a.cell(row=tot, column=c_i, value=f"=SUM({L}4:{L}{tot - 1})").font = BOLD
    for r_i in range(4, tot + 1):
        for c_i in range(1, 7):
            cell = ws_a.cell(row=r_i, column=c_i)
            cell.border = BORDER
            if cell.font != BOLD:
                cell.font = NORMAL
            cell.number_format = NUM if c_i == 6 else INT
    ws_a.cell(row=tot + 2, column=1,
              value="Stock breaks are dated the business date (age 0); cash breaks are aged from the books value date "
                    "(or custodian date if missing in books).").font = Font(name=FONT, italic=True, color="595959")

    # ---------------- Assumptions ----------------
    ws_x.column_dimensions["A"].width = 40
    ws_x.column_dimensions["B"].width = 22
    ws_x.column_dimensions["C"].width = 70
    ws_x["A1"] = "Control rules used in this run (from config.py)"
    ws_x["A1"].font = TITLE
    rows = [
        ("Cash amount tolerance (absolute)", config.AMOUNT_TOLERANCE, "Absorbs paise/cent rounding only."),
        ("Stock quantity tolerance", config.QUANTITY_TOLERANCE, "Positions reconcile to the exact share."),
        ("USD to INR (reporting only)", config.FX_TO_INR["USD"], "Illustrative fixed rate, not market data."),
        ("HIGH priority value (INR eq.)", config.HIGH_VALUE_INR, "Break value at or above this is HIGH."),
        ("MEDIUM priority value (INR eq.)", config.MEDIUM_VALUE_INR, "Break value at or above this is at least MEDIUM."),
        ("HIGH priority age (days)", config.HIGH_AGE_DAYS, "Open this many days or more is HIGH."),
        ("MEDIUM priority age (days)", config.MEDIUM_AGE_DAYS, "Open this many days or more is at least MEDIUM."),
        ("Corporate action ratios checked", ", ".join(str(r) for r in config.CORPORATE_ACTION_RATIOS),
         "Custodian qty / books qty equal to one of these suggests an unbooked split or bonus."),
    ]
    for i, h in enumerate(["Rule", "Value", "Why"], start=1):
        c = ws_x.cell(row=3, column=i, value=h)
        c.font, c.fill, c.border = WHITE_BOLD, HEADER_FILL, BORDER
    for r_i, (a, b, c) in enumerate(rows, start=4):
        for c_i, v in enumerate((a, b, c), start=1):
            cell = ws_x.cell(row=r_i, column=c_i, value=v)
            cell.font, cell.border = NORMAL, BORDER
            if c_i == 2:
                cell.font = Font(name=FONT, color="0000FF")

    wb.save(path)
